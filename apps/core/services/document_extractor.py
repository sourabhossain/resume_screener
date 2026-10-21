"""
Document text extraction from PDF and DOCX files.
"""
import io
import logging
from pathlib import Path
from typing import List

from pypdf import PdfReader
import pypdfium2 as pdfium
import docx

from .llm_client import llm_client

logger = logging.getLogger(__name__)

OCR_SYSTEM_PROMPT = "You are a precise OCR engine for resumes."
OCR_PROMPT = (
    "Transcribe all text in these resume page images exactly as written, in reading order. "
    "Keep section headings, bullet points, dates, emails, phone numbers and URLs. "
    "Output plain text only, with no commentary. "
    "If no text is readable, output nothing."
)


class DocumentExtractor:
    """Extract text content from resume documents."""
    
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx'}

    # Below these, the text is treated as missing. An image-only CV with a stray
    # text overlay still yields a few characters (e.g. "(8th)"), so "not empty"
    # is not enough — screening that text would silently score the candidate ~0.
    MIN_TEXT_CHARS = 100
    MIN_TEXT_WORDS = 20

    # OCR fallback for scanned / image-only PDFs: render pages and read them with
    # the vision-capable LLM. 2x scale = 144 DPI, enough for resume body text.
    # The longest side is capped so oversized pages (posters, odd exports) cannot
    # blow up worker memory; the vision API downsizes past 2048px anyway.
    OCR_MAX_PAGES = 4
    OCR_RENDER_SCALE = 2
    OCR_MAX_IMAGE_SIDE = 2048
    
    @classmethod
    def extract(cls, file_path: str) -> str:
        """
        Extract text from a document file.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Extracted text content
            
        Raises:
            ValueError: If file type is not supported
            FileNotFoundError: If file does not exist
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = path.suffix.lower()
        
        if extension == '.pdf':
            return cls._extract_from_pdf(file_path)
        elif extension == '.docx':
            return cls._extract_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {extension}")
    
    @classmethod
    def has_enough_text(cls, text: str) -> bool:
        """Whether extracted text is substantial enough to screen."""
        stripped = (text or '').strip()
        return len(stripped) >= cls.MIN_TEXT_CHARS and len(stripped.split()) >= cls.MIN_TEXT_WORDS

    @classmethod
    def _extract_from_pdf(cls, file_path: str) -> str:
        """Extract text from PDF file, falling back to OCR for image-only PDFs."""
        text = cls._extract_pdf_text_layer(file_path)
        if cls.has_enough_text(text):
            return text

        logger.info(
            f"PDF has little text ({len(text.strip())} chars), falling back to OCR: {file_path}"
        )
        ocr_text = cls._ocr_pdf(file_path)
        return ocr_text if len(ocr_text.strip()) > len(text.strip()) else text

    @classmethod
    def _ocr_pdf(cls, file_path: str) -> str:
        """Read text from rendered PDF pages with the vision LLM."""
        images = cls._render_pdf_pages(file_path)
        if not images:
            return ""
        return llm_client.invoke_vision(OCR_PROMPT, images, OCR_SYSTEM_PROMPT) or ""

    @classmethod
    def _render_pdf_pages(cls, file_path: str) -> List[bytes]:
        """Render the first OCR_MAX_PAGES pages of a PDF as JPEG bytes."""
        pdf = pdfium.PdfDocument(file_path)
        try:
            images = []
            for index in range(min(len(pdf), cls.OCR_MAX_PAGES)):
                page = pdf[index]
                try:
                    longest_side = max(page.get_size())
                    if longest_side <= 0:
                        continue
                    scale = min(cls.OCR_RENDER_SCALE, cls.OCR_MAX_IMAGE_SIDE / longest_side)
                    image = page.render(scale=scale).to_pil().convert('RGB')
                finally:
                    page.close()
                buffer = io.BytesIO()
                image.save(buffer, format='JPEG', quality=85)
                images.append(buffer.getvalue())
            return images
        finally:
            pdf.close()

    @staticmethod
    def _extract_pdf_text_layer(file_path: str) -> str:
        """Extract the embedded text layer from a PDF file."""
        text_parts = []
        
        with open(file_path, 'rb') as file:
            reader = PdfReader(file)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        
        return "\n".join(text_parts)
    
    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        """Extract text from DOCX file."""
        doc = docx.Document(file_path)
        text_parts = []
        

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        
        return "\n".join(text_parts)
    
    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        """Check if file type is supported."""
        extension = Path(file_path).suffix.lower()
        return extension in cls.SUPPORTED_EXTENSIONS
